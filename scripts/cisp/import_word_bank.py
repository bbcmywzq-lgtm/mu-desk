from __future__ import annotations

import argparse
import io
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from PIL import Image

from build_expanded_bank import (
    Question,
    find_duplicate,
    fingerprint,
    load_existing_questions,
    normalize,
    option_similarity,
    write_markdown,
)


ANSWER_RE = re.compile(
    r"^\s*(?:正确答案|参考答案|答案)\s*(?:是|为)?\s*[：:]?\s*[（(]?\s*([A-DＡ-Ｄ])\s*[）)]?",
    re.I | re.M,
)
EXPLANATION_RE = re.compile(r"^\s*(?:解析|解释|注)\s*[：:]\s*", re.I)
OPTION_START_RE = re.compile(
    r"^\s*([A-D])(?:\s*[.,，、．。:：]\s*|\s+|(?=[\u4e00-\u9fff0-9《（]))(.*)$",
    re.I,
)
INLINE_MARKER_RE = re.compile(
    r"(?<!\S)([A-D])(?:\s*[.,，、．。:：]\s*|\s+|(?=[\u4e00-\u9fff0-9《（]))",
    re.I,
)
HEADER_RE = re.compile(
    r"^(?:CERTIFICATION\b|重要练习题|注册信息安全专业人员|大纲知识点|知识点复习|模拟考试|"
    r"（?时间[：:]|姓名\s+单位名称|（极为重要)",
    re.I,
)


@dataclass
class ParagraphRecord:
    text: str
    image_names: list[str]


@dataclass
class ParsedDocument:
    questions: list[Question]
    diagnostics: list[str]
    answer_markers: int


def clean_line(value: str) -> str:
    value = value.translate(str.maketrans("ＡＢＣＤ．，：　", "ABCD.,: "))
    value = value.replace("", " ").replace("", " ").replace("", " ")
    return normalize(value)


def repair_option_boundaries(value: str) -> str:
    compact_markers = re.findall(
        r"(?<=[\u4e00-\u9fff0-9）)])([B-D])(?=[.,，、．。:：\u4e00-\u9fff0-9《（])",
        value,
    )
    if len(set(compact_markers)) >= 2:
        value = re.sub(
            r"(?<=[\u4e00-\u9fff0-9）)])([B-D])(?=[.,，、．。:：\u4e00-\u9fff0-9《（])",
            r" \1",
            value,
        )
    return value


def numeric_signature(value: str) -> list[str]:
    return sorted(re.findall(r"(?<![A-Za-z])\d+(?:\.\d+)?%?(?![A-Za-z])", normalize(value)))


def relationship_ids(paragraph) -> list[str]:
    result: list[str] = []
    for element in paragraph._p.iter():
        local_name = element.tag.rsplit("}", 1)[-1]
        if local_name not in {"blip", "imagedata"}:
            continue
        for key, value in element.attrib.items():
            if key.rsplit("}", 1)[-1] in {"embed", "id"} and value not in result:
                result.append(value)
    return result


def save_image(blob: bytes, path: Path) -> None:
    with Image.open(io.BytesIO(blob)) as image:
        image.load()
        if image.mode in {"RGBA", "LA"}:
            converted = image.convert("RGBA")
        else:
            converted = image.convert("RGB")
        path.parent.mkdir(parents=True, exist_ok=True)
        converted.save(path, format="PNG", optimize=True)


def paragraph_records(document, image_dir: Path, document_index: int) -> list[ParagraphRecord]:
    relationship_names: dict[str, str] = {}
    records: list[ParagraphRecord] = []
    for paragraph in document.paragraphs:
        image_names: list[str] = []
        for relationship_id in relationship_ids(paragraph):
            relationship = document.part.rels.get(relationship_id)
            if relationship is None or "image" not in relationship.reltype:
                continue
            if relationship_id not in relationship_names:
                image_name = f"wx2026-{document_index:02d}-{len(relationship_names) + 1:02d}.png"
                try:
                    save_image(relationship.target_part.blob, image_dir / image_name)
                except Exception:
                    continue
                relationship_names[relationship_id] = image_name
            image_names.append(relationship_names[relationship_id])
        records.append(ParagraphRecord(paragraph.text, image_names))
    return records


def last_question_block(records: list[ParagraphRecord], start: int, end: int) -> list[ParagraphRecord]:
    cursor = end - 1
    while cursor >= start and not records[cursor].text.strip() and not records[cursor].image_names:
        cursor -= 1
    block_end = cursor + 1
    while cursor >= start and (records[cursor].text.strip() or records[cursor].image_names):
        cursor -= 1
    block_start = cursor + 1
    block = records[block_start:block_end]
    first_text = clean_line(next((record.text for record in block if record.text.strip()), ""))
    if OPTION_START_RE.match(first_text):
        # Some Word files put a blank paragraph between the stem/figure and the
        # options.  Pull in the immediately preceding text block when the last
        # block begins directly with A-D choices.
        while cursor >= start and not records[cursor].text.strip() and not records[cursor].image_names:
            cursor -= 1
        previous_end = cursor + 1
        while cursor >= start and (records[cursor].text.strip() or records[cursor].image_names):
            cursor -= 1
        previous = records[cursor + 1 : previous_end]
        previous_text = clean_line(" ".join(record.text for record in previous))
        if previous_text and not EXPLANATION_RE.match(previous_text):
            block = previous + block
    return block


def split_inline_options(line: str) -> tuple[list[str], list[str]] | None:
    matches = list(INLINE_MARKER_RE.finditer(line))
    results: list[tuple[list[str], list[str]]] = []
    for start_index, match in enumerate(matches):
        if match.group(1).upper() != "A":
            continue
        chosen = [match]
        expected = "B"
        for candidate in matches[start_index + 1 :]:
            if candidate.group(1).upper() == expected:
                chosen.append(candidate)
                if expected == "D":
                    break
                expected = chr(ord(expected) + 1)
        if len(chosen) != 4:
            continue
        prefix = clean_line(line[: chosen[0].start()])
        options: list[str] = []
        for option_index, marker in enumerate(chosen):
            option_end = chosen[option_index + 1].start() if option_index + 1 < 4 else len(line)
            options.append(clean_line(line[marker.end() : option_end]))
        if all(options):
            results.append((([prefix] if prefix else []), options))
    # If the stem itself mentions subjects A/B, the actual A-D option sequence
    # is the later one in the paragraph.
    return results[-1] if results else None


def split_option_fragments(line: str, expected_index: int) -> tuple[str, list[str]] | None:
    matches = list(INLINE_MARKER_RE.finditer(line))
    expected_letter = chr(ord("A") + expected_index)
    candidates: list[tuple[int, int, str, list[str]]] = []
    for start_index, match in enumerate(matches):
        if match.group(1).upper() != expected_letter:
            continue
        chosen = [match]
        next_letter = chr(ord(expected_letter) + 1) if expected_letter != "D" else None
        for candidate in matches[start_index + 1 :]:
            if next_letter is None or candidate.group(1).upper() != next_letter:
                break
            chosen.append(candidate)
            next_letter = chr(ord(next_letter) + 1) if next_letter != "D" else None
        values: list[str] = []
        for option_index, marker in enumerate(chosen):
            option_end = chosen[option_index + 1].start() if option_index + 1 < len(chosen) else len(line)
            values.append(clean_line(line[marker.end() : option_end]))
        if all(values) and (len(chosen) > 1 or match.start() == 0):
            candidates.append((len(chosen), match.start(), clean_line(line[: match.start()]), values))
    if expected_index < 3:
        next_letter = chr(ord("A") + expected_index + 1)
        for start_index, match in enumerate(matches):
            if match.group(1).upper() != next_letter:
                continue
            prefix_value = clean_line(line[: match.start()])
            if not prefix_value:
                continue
            chosen = [match]
            wanted = chr(ord(next_letter) + 1) if next_letter != "D" else None
            for candidate in matches[start_index + 1 :]:
                if wanted is None or candidate.group(1).upper() != wanted:
                    break
                chosen.append(candidate)
                wanted = chr(ord(wanted) + 1) if wanted != "D" else None
            values = [prefix_value]
            for option_index, marker in enumerate(chosen):
                option_end = chosen[option_index + 1].start() if option_index + 1 < len(chosen) else len(line)
                values.append(clean_line(line[marker.end() : option_end]))
            if all(values):
                candidates.append((len(values), match.start(), "", values))
    if not candidates:
        return None
    _, _, prefix, values = max(candidates, key=lambda item: (item[0], item[1]))
    return prefix, values


def parse_question_block(block: list[ParagraphRecord], answer_letter: str, source_number: int) -> Question | None:
    lines: list[str] = []
    image_names: list[str] = []
    for record in block:
        image_names.extend(name for name in record.image_names if name not in image_names)
        lines.extend(repair_option_boundaries(clean_line(line)) for line in record.text.splitlines() if clean_line(line))
    lines = [line for line in lines if not HEADER_RE.search(line)]
    while lines and EXPLANATION_RE.match(lines[0]):
        lines.pop(0)
    if not lines:
        return None

    stem_lines: list[str] = []
    options: list[str] = []
    current_option: int | None = None
    option_started = False
    for line in lines:
        fragments = split_option_fragments(line, len(options)) if len(options) < 4 else None
        if fragments is not None:
            inline_prefix, inline_options = fragments
            if not option_started and inline_prefix:
                stem_lines.append(inline_prefix)
            options.extend(inline_options)
            option_started = True
            current_option = len(options) - 1
            continue
        match = OPTION_START_RE.match(line)
        if match:
            letter_index = ord(match.group(1).upper()) - ord("A")
            if letter_index == len(options) and letter_index < 4:
                options.append(clean_line(match.group(2)))
                option_started = True
                current_option = letter_index
                continue
        if option_started and current_option is not None:
            options[current_option] = clean_line(options[current_option] + " " + line)
        else:
            stem_lines.append(line)

    if len(options) != 4 or any(not option for option in options):
        # Several reviewed sheets omit A-D labels entirely.  In those blocks,
        # the final four paragraphs are the choices and everything before them
        # is the question/context.
        if len(lines) >= 5:
            fallback_options = [clean_line(value) for value in lines[-4:]]
            fallback_stem = lines[:-4]
            if all(fallback_options) and fallback_stem:
                stem_lines, options = fallback_stem, fallback_options

    if (len(options) != 4 or any(not option for option in options)) and image_names:
        options = [f"选项 {letter}（见题图）" for letter in "ABCD"]
    if len(options) != 4 or any(not option for option in options):
        return None

    cleaned_options: list[str] = []
    for option_index, value in enumerate(options):
        letter = chr(ord("A") + option_index)
        value = re.sub(rf"^{letter}\s*[.,，、．。:：]\s*", "", value, flags=re.I)
        value = re.sub(rf"^{letter}\s+", "", value, flags=re.I)
        cleaned_options.append(clean_line(value))
    options = cleaned_options
    if any(not option for option in options):
        return None

    stem = clean_line(" ".join(stem_lines))
    stem = re.sub(r"^\s*\d{1,3}\s*[.,，、．:：]\s*", "", stem)
    stem = re.sub(rf"^\s*{source_number}\s*(?:[.,，、．]\s*|(?=[\u4e00-\u9fff]))", "", stem)
    # Some reviewed sheets print the key in parentheses in the stem.  Keep the
    # blank but remove the answer leak before importing it into practice mode.
    stem = re.sub(r"[（(]\s*[A-D]\s*[）)]", "（ ）", stem, flags=re.I)
    stem = clean_line(stem)
    if len(stem) < 6 or "答案" in stem:
        return None
    return Question(
        stem=stem,
        options=options,
        correct_index=ord(answer_letter.upper()) - ord("A"),
        image_names=image_names,
    )


def explanation_after(records: list[ParagraphRecord], answer_index: int, inline_suffix: str = "") -> str:
    parts: list[str] = []
    suffix = clean_line(inline_suffix)
    if EXPLANATION_RE.search(suffix):
        parts.append(EXPLANATION_RE.sub("", suffix))
    cursor = answer_index + 1
    while cursor < len(records):
        text = clean_line(records[cursor].text)
        if not text:
            if parts:
                break
            cursor += 1
            continue
        if not parts and not EXPLANATION_RE.search(text):
            break
        parts.append(EXPLANATION_RE.sub("", text))
        cursor += 1
    return clean_line(" ".join(parts)) or "公开题库未附解析。"


def parse_document(path: Path, image_dir: Path, document_index: int) -> ParsedDocument:
    document = Document(path)
    records = paragraph_records(document, image_dir, document_index)
    answer_markers: list[tuple[int, str, str, str]] = []
    for index, record in enumerate(records):
        searchable = record.text.translate(str.maketrans("ＡＢＣＤ．，：　", "ABCD.,: "))
        match = ANSWER_RE.search(searchable)
        if match:
            answer_markers.append((index, clean_line(match.group(1)), searchable[: match.start()], searchable[match.end() :]))

    questions: list[Question] = []
    diagnostics: list[str] = []
    previous_answer = -1
    for source_number, (answer_index, answer_letter, inline_prefix, inline_suffix) in enumerate(answer_markers, start=1):
        block = last_question_block(records, previous_answer + 1, answer_index)
        if clean_line(inline_prefix):
            block.append(ParagraphRecord(inline_prefix, records[answer_index].image_names))
        question = parse_question_block(block, answer_letter, source_number)
        if question is None:
            preview = clean_line(" ".join(record.text for record in block))[:100]
            diagnostics.append(f"{path.name} marker {source_number}: could not parse {preview!r}")
        else:
            question.explanation = explanation_after(records, answer_index, inline_suffix)
            question.origin = f"wx2026-{document_index:02d}"
            question.origin_number = str(source_number)
            questions.append(question)
        previous_answer = answer_index
    return ParsedDocument(questions, diagnostics, len(answer_markers))


def merge(base: list[Question], candidates: list[Question], report: list[str]) -> tuple[int, int]:
    exact = {fingerprint(question.stem): index for index, question in enumerate(base)}
    added = 0
    duplicates = 0
    for candidate in candidates:
        duplicate_index = find_duplicate(candidate, base, exact)
        if duplicate_index is None:
            candidate_key = fingerprint(candidate.stem)
            candidate_numbers = numeric_signature(candidate.stem)
            calculation_sensitive = bool(
                re.search(r"计算(?!机)|假设|损失|价值|缺陷密度|年度发生率|RTO|RPO", candidate.stem, re.I)
            )
            best: tuple[float, float, int] | None = None
            for index, existing in enumerate(base):
                stem_score = SequenceMatcher(None, candidate_key, fingerprint(existing.stem)).ratio()
                same_numbers = candidate_numbers == numeric_signature(existing.stem)
                options_score = option_similarity(candidate, existing) if stem_score >= 0.70 else 0.0
                same_question_with_changed_values = stem_score >= 0.84 and options_score >= 0.98
                same_question_with_ocr_noise = stem_score >= 0.79 and options_score >= 0.84
                if not same_numbers and not (
                    same_question_with_changed_values
                    or same_question_with_ocr_noise
                    or (stem_score >= 0.95 and options_score >= 0.90 and not calculation_sensitive)
                ):
                    continue
                if best is None or stem_score > best[0]:
                    best = (stem_score, options_score, index)
            if best is not None:
                stem_score, options_score, best_index = best
                both_figure_questions = "图" in candidate.stem and "图" in base[best_index].stem
                if (
                    stem_score >= 0.94
                    or (stem_score >= 0.92 and min(len(candidate_key), len(fingerprint(base[best_index].stem))) < 40)
                    or (stem_score >= 0.87 and options_score >= 0.88)
                    or (stem_score >= 0.82 and options_score >= 0.95)
                    or (stem_score >= 0.84 and options_score >= 0.98)
                    or (stem_score >= 0.79 and options_score >= 0.84)
                    or (both_figure_questions and stem_score >= 0.82)
                ):
                    duplicate_index = best_index
        if duplicate_index is not None:
            duplicates += 1
            existing = base[duplicate_index]
            if existing.explanation == "公开题库未附解析。" and candidate.explanation != "公开题库未附解析。":
                existing.explanation = candidate.explanation
            if not existing.image_names and candidate.image_names:
                existing.image_names = list(candidate.image_names)
            continue
        base.append(candidate)
        exact[fingerprint(candidate.stem)] = len(base) - 1
        added += 1
    return added, duplicates


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--converted-dir", type=Path, required=True)
    parser.add_argument("--base-questions-js", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--image-dir", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()

    source_files = sorted(args.source_dir.glob("*.docx"))
    converted_files = sorted(args.converted_dir.glob("*.docx"))
    files = source_files + converted_files
    if len(files) != 8:
        raise ValueError(f"Expected eight Word files, got {len(files)}")

    base = load_existing_questions(args.base_questions_js)
    original_count = len(base)
    original_fingerprint_counts = Counter(fingerprint(question.stem) for question in base)
    original_duplicate_keys = {key for key, count in original_fingerprint_counts.items() if key and count > 1}
    report = [f"Preserved baseline: {original_count}"]
    all_candidates: list[Question] = []
    for document_index, path in enumerate(files, start=1):
        parsed = parse_document(path, args.image_dir, document_index)
        all_candidates.extend(parsed.questions)
        report.append(
            f"{document_index:02d} {path.name}: markers={parsed.answer_markers}, parsed={len(parsed.questions)}, skipped={len(parsed.diagnostics)}"
        )
        report.extend(f"SKIP {value}" for value in parsed.diagnostics)

    added, duplicates = merge(base, all_candidates, report)
    report.append(f"Merge: candidates={len(all_candidates)}, added={added}, duplicates={duplicates}, total={len(base)}")

    fingerprints = [fingerprint(question.stem) for question in base]
    exact_duplicates = [key for key, count in Counter(fingerprints).items() if key and count > 1]
    introduced_duplicates = [key for key in exact_duplicates if key not in original_duplicate_keys]
    if introduced_duplicates:
        raise ValueError(f"Exact duplicate audit failed: import introduced {len(introduced_duplicates)} duplicate stems")
    report.append(f"Pre-existing exact duplicate stems preserved for progress compatibility: {len(original_duplicate_keys)}")

    referenced = {name for question in base for name in question.image_names if name.startswith("wx2026-")}
    missing_images = [name for name in referenced if not (args.image_dir / name).is_file()]
    if missing_images:
        raise ValueError(f"Missing imported images: {missing_images}")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    write_markdown(args.output, base)
    report.append(f"Referenced imported images: {len(referenced)}")
    args.report.write_text("\n".join(report) + "\n", encoding="utf-8")
    print("\n".join(report[-12:]))
    print(f"FINAL_QUESTION_COUNT={len(base)}")
    return 0 if len(base) > original_count else 4


if __name__ == "__main__":
    sys.exit(main())
